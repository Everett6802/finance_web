#! /usr/bin/python
# -*- coding: utf8 -*-

import os
import sys
import re
import errno
# # '''
# # Question: How to Solve xlrd.biffh.XLRDError: Excel xlsx file; not supported ?
# # Answer : The latest version of xlrd(2.01) only supports .xls files. Installing the older version 1.2.0 to open .xlsx files.
# # '''
# import xlrd
# # import xlsxwriter
from openpyxl import Workbook, load_workbook
import argparse
from datetime import datetime, date, timedelta
import getpass
from collections import OrderedDict
import statistics


DISTRIBUTION_STATISTICS = {
    '0': 'all',
    '1': 'change_pct',
    '2': 'volume',
    'all': 'all',
    'change_pct': 'change_pct',
    'volume': 'volume',
}

def parse_distribution_statistics(value):
    if value not in DISTRIBUTION_STATISTICS:
        raise argparse.ArgumentTypeError('Invalid distribution statistics. Choices: 0/all, 1/change_pct, 2/volume')
    return DISTRIBUTION_STATISTICS[value]


class RollingStatistics(object):

	DEFAULT_HOST_DATA_FOLDERPATH =  "C:\\Users\\%s\\project_data\\finance_web" % getpass.getuser()
	DEFAULT_DATA_FOLDERPATH =  os.getenv("DATA_PATH", DEFAULT_HOST_DATA_FOLDERPATH)
	DEFAULT_SOURCE_FILENAME = "加權指數歷史資料2000-2025.xlsx"
	DEFAULT_TIME_FIELD_NAME = "時間"	
	# DEFAULT_CLOSING_PRICE_FIELD_NAME = "收盤價"	
	DEFAULT_VOLUME_FIELD_NAME = "成交量"
	DEFAULT_EXT_CHANGE_PCT_FIELD_TITLE = "漲跌幅"
	# DEFAULT_CONFIG_FOLDERPATH =  "C:\\Users\\%s" % os.getlogin()
	DEFAULT_DATE_BASE_NUMBER = 36526
	DEFAULT_DATE_BASE = date(2000, 1, 1)
	DEFAULT_ROLLING_WINDOW_SIZE = 240
	DEFAULT_DATE_FORMAT = "%Y-%m-%d"
	DEFAULT_DISTRIBUTION_STATISTICS = "all"

	@classmethod
	def __is_string(cls, value):
		is_string = False
		try:
			int(value)
		except ValueError:
			is_string = True
		return is_string


	@classmethod
	def __check_file_exist(cls, filepath):
		check_exist = True
		try:
			os.stat(filepath)
		except OSError as exception:
			if exception.errno != errno.ENOENT:
				print ("%s: %s" % (errno.errorcode[exception.errno], os.strerror(exception.errno)))
				raise
			check_exist = False
		return check_exist


	@classmethod
	def __is_excel_locked(cls, filepath):
# Excel lock file
		lock_path = os.path.join(os.path.dirname(filepath), "~$" + os.path.basename(filepath))
		if os.path.exists(lock_path):
			return True
# OS lock
		try:
			with open(filepath, "a"):
				return False
		except PermissionError:
			return True


# ========= XLSX 工具 =========
	@classmethod
	def __read_xlsx(cls, source_filepath, rolling_window_size=None, sheet_name=None):
		wb = load_workbook(source_filepath, read_only=True, data_only=True)
		ws = wb.active if sheet_name is None else wb[sheet_name]
		rows = []
		header = [c.value for c in ws[1]]
		rows.append(header)
		# for row in ws.iter_rows(min_row=2, values_only=True):
		# 	rows.append(dict(zip(header, row)))
		start_row = 2 
		if rolling_window_size is not None:
			start_row = max(2, ws.max_row - rolling_window_size + 1)
		for row in ws.iter_rows(min_row=start_row, values_only=True):
			rows.append(row)
		return rows


	@classmethod
	def __get_percentiles(cls, data, percentiles, descending=False):
		sorted_data = sorted(data, reverse=descending)
		n = len(sorted_data)
		if n == 0:
			return None
		percentile_values = []
		for percentile in percentiles:
			pos = (n - 1) * percentile / 100
			lower = int(pos)
			upper = min(lower + 1, n - 1)
			weight = pos - lower
			percentile_values.append(sorted_data[lower] * (1 - weight) + sorted_data[upper] * weight)
		return percentile_values


	@classmethod
	def __get_percentile(cls, data, percentile, descending=False):
		percentiles = [percentile]
		percentile_values = cls.__get_percentiles(data, percentiles, descending=descending)
		return percentile_values[0] if percentile_values is not None else None
		# sorted_data = sorted(data, reverse=descending)
		# n = len(sorted_data)
		# if n == 0:
		# 	return None
		# pos = (n - 1) * percentile / 100
		# lower = int(pos)
		# upper = min(lower + 1, n - 1)
		# weight = pos - lower
		# return sorted_data[lower] * (1 - weight) + sorted_data[upper] * weight


	@classmethod
	def __get_percentile_rank(cls, data, value, descending=False):
		sorted_data = sorted(data, reverse=descending)
		n = len(sorted_data)
		if n == 0:
			return None
		func_ptr = lambda x: x <= value if not descending else x >= value
		# filtered_values = list(filter(func_ptr, sorted_data))
		# return len(filtered_values) / n * 100
		count = sum(1 for x in sorted_data if func_ptr(x))
		return count / n * 100


	@classmethod
	def __get_statistics(cls, data, descending=False):
		if len(data) == 0:
			return None
		percentile_ranks = [1, 5, 10, 25, 50, 75, 90, 95, 99]
		percentiles = cls.__get_percentiles(data, percentile_ranks, descending=descending)
		percentiles_dict = {"p%d" % k: v for k, v in zip(percentile_ranks, percentiles)}
		return {
			"count": len(data),
			"mean": statistics.mean(data),
# statistics.stdev(data) 在只有 1 筆資料時會報錯 => 我個人反而比較喜歡不要偷偷把它變成 0，因為「只有一筆資料，標準差不存在」
			"stdev": statistics.stdev(data) if len(data) > 1 else None,
			"min": min(data),
			"max": max(data),
			**percentiles_dict
		}


	def __init__(self, cfg):
		self.xcfg = {
			"source_folderpath": None,
			# "source_filename": None,
			"stock_symbol_string": None,
			"check_date": None,
			"target_change_pct_value": None,
			"target_volume_value": None,
			"risk_free_rate": 0.0,
			"statistics_date_range_string": None,
			"rolling_window_size": None,
			"start_date": None,
			"end_date": None,
			"distribution_statistics": self.DEFAULT_DISTRIBUTION_STATISTICS,
		}
		# import pdb; pdb.set_trace()
		self.xcfg.update(cfg)
		self.xcfg["source_folderpath"] = self.DEFAULT_DATA_FOLDERPATH if self.xcfg["source_folderpath"] is None else self.xcfg["source_folderpath"]
		# self.xcfg["source_filename"] = self.DEFAULT_SOURCE_FILENAME if self.xcfg["source_filename"] is None else self.xcfg["source_filename"]
		self.cur_year = datetime.now().year
		self.filepath_dict = OrderedDict()
		self.filepath_dict["source_folderpath"] = self.xcfg["source_folderpath"]
		self.statistics_by_date_range = True 
		if self.xcfg["start_date"] is None and self.xcfg["end_date"] is None:
			self.statistics_by_date_range = False
			if self.xcfg["rolling_window_size"] is None:
				self.xcfg["rolling_window_size"] = self.DEFAULT_ROLLING_WINDOW_SIZE


	def __enter__(self):
		return self


	def __exit__(self, type, msg, traceback):
		return False


	def __read_worksheet(self, source_filepath, expected_title_list=None):
# Check if it's required to transform from stock name to stock symbol
		worksheet_data = {}
		# import pdb; pdb.set_trace()
		worksheet_rows = self.__read_xlsx(source_filepath, self.xcfg["rolling_window_size"])
		sheet_ncols = len(worksheet_rows[0])
		sheet_nrows = len(worksheet_rows)
		data_list = []
# title
		title_list = []
		for column_index in range(0, sheet_ncols):
			title_value = worksheet_rows[0][column_index]
			title_list.append(title_value)
		# print(title_list)
		# import pdb; pdb.set_trace()
		time_index = None
		title_index_list = None
		if expected_title_list is not None:
			if self.statistics_by_date_range:
				time_index = expected_title_list.index(self.DEFAULT_TIME_FIELD_NAME)
			# time_index = title_list.index(self.DEFAULT_TIME_FIELD_NAME)
			# title_index_list = [time_index,]
			title_index_list = []
			for expected_title in expected_title_list:
				index = title_list.index(expected_title)
				title_index_list.append(index)
			new_title_list = [title for index, title in enumerate(title_list) if index in title_index_list]
			title_list = new_title_list
		else:
			if self.statistics_by_date_range:
				time_index = title_list.index(self.DEFAULT_TIME_FIELD_NAME)
			# title_index_list = list(range(0, self.Worksheet.ncols))
			title_index_list = list(range(0, sheet_ncols))

		start_date_obj = None
		end_date_obj = None
		date_range_funcptr = None
		if self.statistics_by_date_range:
			if time_index is None:
				raise ValueError("The title list must include the time field name: %s" % self.DEFAULT_TIME_FIELD_NAME)
			if self.xcfg["start_date"] is not None:
				start_date_obj = datetime.strptime(self.xcfg["start_date"], self.DEFAULT_DATE_FORMAT)
			if self.xcfg["end_date"] is not None:
				end_date_obj = datetime.strptime(self.xcfg["end_date"], self.DEFAULT_DATE_FORMAT)
			if self.xcfg["start_date"] is not None and self.xcfg["end_date"] is not None:
				date_range_funcptr = lambda x: True if start_date_obj <= datetime.strptime(x, self.DEFAULT_DATE_FORMAT) <= end_date_obj else False
			elif self.xcfg["start_date"] is not None:
				date_range_funcptr = lambda x: True if start_date_obj <= datetime.strptime(x, self.DEFAULT_DATE_FORMAT) else False
			elif self.xcfg["end_date"] is not None:
				date_range_funcptr = lambda x: True if datetime.strptime(x, self.DEFAULT_DATE_FORMAT) <= end_date_obj else False
# data
		# for row_index in range(1, self.Worksheet.nrows):
		for row_index in range(1, sheet_nrows):
			entry_list = []
			# can_add = True
			# for column_index in range(0, self.Worksheet.ncols):
			for column_index in title_index_list:
				# entry_value = self.Worksheet.cell_value(row_index, column_index)
				entry_value = worksheet_rows[row_index][column_index]
				# if column_index == time_index:
				# 	day_diff = int(entry_value) - self.DEFAULT_DATE_BASE_NUMBER
				# 	entry_date = self.DEFAULT_DATE_BASE + timedelta(days=day_diff)
				# 	# if entry_date < date(2010, 1, 1):
				# 	# 	can_add = False
				# 	# 	break
				# 	entry_value = entry_date.strftime("%m/%d/%Y")
				# 	# print("%d: %s" % (row_index + 1, entry_value))
				# 	# import pdb; pdb.set_trace()
				entry_list.append(entry_value)
			# print("%d: %s" % (row_index + 1, entry_list))
			# if can_add: data_list.append(entry_list)
			if self.statistics_by_date_range:
				if not date_range_funcptr(entry_list[time_index]):
					continue
			data_list.append(entry_list)
		worksheet_data["title"] = title_list
		worksheet_data["data"] = data_list
		return worksheet_data


	def __date_str2list(self, date_str, skip_year=False):
		# import pdb; pdb.set_trace()
		# print(date_str)
		if date_str.find("/") != -1:
			elem_list = date_str.split("/")
			if len(elem_list) == 2:
				[year, month, day] = [self.cur_year, int(elem_list[0]), int(elem_list[1])]
			elif len(elem_list) == 3:
				[month, day, year] = list(map(int, elem_list))
			else:
				raise ValueError("Incorrect date string format: %s" % date_str)
		elif date_str.find("-") != -1:
			elem_list = date_str.split("-")
			if len(elem_list) == 2:
				[year, month, day] = [self.cur_year, int(elem_list[0]), int(elem_list[1])]
			elif len(elem_list) == 3:
				[year, month, day] = list(map(int, elem_list))
			else:
				raise ValueError("Incorrect date string format: %s" % date_str)
		else:
			raise ValueError("Incorrect date string format: %s" % date_str)
		return [month, day] if skip_year else [year, month, day]


	def __date_str2obj(self, date_str):
		[year, month, day] = self.__date_str2list(date_str)
		date_obj = None
		try:
			date_obj = date(year, month, day)
		except Exception as e:
			print("Unsupport date format[%s] due to: %s" % (date_str, str(e)))
			# import pdb; pdb.set_trace()
			raise e
		return date_obj


	def __date_number2obj(self, date_number):
		day_diff = int(date_number) - self.DEFAULT_DATE_BASE_NUMBER
		date_obj = self.DEFAULT_DATE_BASE + timedelta(days=day_diff)
		return date_obj


	def __date_number2list(self, date_number):
		date_obj = self.__date_number2obj(date_number)
		return [date_obj.year, date_obj.month, date_obj.day]


	# def __extract_data(self, source_filename, rolling_window_size=None):
	# 	# import pdb; pdb.set_trace()
	# 	expected_title_list = [self.DEFAULT_TIME_FIELD_NAME, self.DEFAULT_CLOSING_PRICE_FIELD_NAME,]
	# 	# worksheet_data = self.__get_worksheet_data(expected_title_list)
	# 	worksheet_data = self.__read_worksheet(source_filename, rolling_window_size, expected_title_list)
	# 	time_index = worksheet_data["title"].index(self.DEFAULT_TIME_FIELD_NAME)
	# 	need_check_time_range = (date_range_start is not None) or (date_range_end is not None)
	# 	date_range_start_list = None
	# 	date_range_end_list = None
	# 	if need_check_time_range:
	# 		if date_range_start is not None:
	# 			date_range_start_list = self.__date_str2list(date_range_start)
	# 		else:
	# 			date_range_start_list = self.__date_str2list("01/01/1900")
	# 		if date_range_end is not None:
	# 			date_range_end_list = self.__date_str2list(date_range_end)
	# 		else:
	# 			date_range_end_list = self.__date_str2list("12/31/2100")
	# 		funcptr = lambda x: date_range_start_list <= self.__date_str2list(x[time_index]) <= date_range_end_list
	# 		data_tmp = filter(funcptr, worksheet_data["data"])
	# 		worksheet_data["data"] = list(data_tmp)
	# 	return worksheet_data


	def print_filepath(self):
		print("************** File Path **************")
		for key, value in self.filepath_dict.items():
			print("%s: %s" % (key, value))


	def __analyze_percentile(self, worksheet_data, expected_title_list, field_name, check_date=None, check_value=None, descending_funcptr=None):
		if self.DEFAULT_TIME_FIELD_NAME not in expected_title_list:
			raise ValueError("The expected title list must include the time field name: %s" % self.DEFAULT_TIME_FIELD_NAME)
		if field_name not in expected_title_list:
			raise ValueError("The expected title list must include the field name: %s" % field_name)
		time_data_index = expected_title_list.index(self.DEFAULT_TIME_FIELD_NAME)
		field_data_index = expected_title_list.index(field_name)
		# import pdb; pdb.set_trace()
		# check_date = None
		# check_value = None
		if check_value is None and check_date is None:
			check_date = worksheet_data["data"][-1][time_data_index]
			check_value = worksheet_data["data"][-1][field_data_index]
		elif check_date is not None:
			if check_value is not None:
				print("Warning: Both check_date and check_value are specified, so the check_value will be ignored")
			try:
				time_list = list(map(lambda x: x[time_data_index], worksheet_data["data"]))
				time_index = time_list.index(check_date)
				check_date = worksheet_data["data"][time_index][time_data_index]
				check_value = worksheet_data["data"][time_index][field_data_index]
			except ValueError as e:
				print("Error: The check date is not found in the data: %s" % check_date)
				raise e
		descending = False
		if descending_funcptr is not None:
			descending = descending_funcptr(check_value)		
		field_data_list = list(map(lambda x: x[field_data_index], worksheet_data["data"]))
		percentile_rank = self.__get_percentile_rank(field_data_list, check_value, descending=descending)
		return check_date, check_value, percentile_rank


	def analyze_change_pct_percentile(self, worksheet_data, expected_title_list):
		descending_funcptr = lambda x: True if x >= 0 else False
		return self.__analyze_percentile(worksheet_data, expected_title_list, self.DEFAULT_EXT_CHANGE_PCT_FIELD_TITLE, check_date=self.xcfg["check_date"], check_value=self.xcfg["target_change_pct_value"], descending_funcptr=descending_funcptr)


	def analyze_volume_percentile(self, worksheet_data, expected_title_list):
		return self.__analyze_percentile(worksheet_data, expected_title_list, self.DEFAULT_VOLUME_FIELD_NAME, check_date=self.xcfg["check_date"], check_value=self.xcfg["target_volume_value"], descending_funcptr=lambda x: True)


	def __show_analysis(self):
		expected_title_list = [self.DEFAULT_TIME_FIELD_NAME, self.DEFAULT_VOLUME_FIELD_NAME, self.DEFAULT_EXT_CHANGE_PCT_FIELD_TITLE,]
		stock_symbol_list = self.xcfg["stock_symbol_string"].split(",")
		check_date = self.xcfg["check_date"]
		for stock_symbol in stock_symbol_list:
			source_filepath = os.path.join(self.xcfg["source_folderpath"], f"{stock_symbol}.xlsx")
			file_exist = self.__check_file_exist(source_filepath)
			if not file_exist:
				print(f"Warning: The file {source_filepath} is not found, so skip fetching data for {stock_symbol}...")
				continue
			if self.__is_excel_locked(source_filepath):
				print(f"Warning: The file {source_filepath} is locked by other process, so skip fetching data for {stock_symbol}...")
				continue
			# import pdb; pdb.set_trace()
			worksheet_data = self.__read_worksheet(source_filepath, expected_title_list)
			print(f"==========  {stock_symbol}  ==========")
			check_change_pct_date, check_change_pct_value, percentile_rank = self.analyze_change_pct_percentile(worksheet_data, expected_title_list)
			print("Change Percentage: date=%s, value=%.2f%%, percentile_rank=%.2f%%" % (check_change_pct_date, check_change_pct_value * 100.0, percentile_rank))
			check_volume_date, check_volume_value, percentile_rank = self.analyze_volume_percentile(worksheet_data, expected_title_list)
			print("Volume: date=%s, value=%d, percentile_rank=%.2f%%" % (check_volume_date, int(check_volume_value / 1000) , percentile_rank))
			print(f"=====================================\n")


	def __show_target_analysis(self):
		if self.xcfg["target_change_pct_value"] is None and self.xcfg["target_volume_value"] is None:
			print("Warning: No target value to check...")
			return
		expected_title_list = [self.DEFAULT_TIME_FIELD_NAME,]
		stock_symbol = self.xcfg["stock_symbol_string"]
		if self.xcfg["target_volume_value"] is not None:
			expected_title_list.append(self.DEFAULT_VOLUME_FIELD_NAME)
		if self.xcfg["target_change_pct_value"] is not None:
			expected_title_list.append(self.DEFAULT_EXT_CHANGE_PCT_FIELD_TITLE)
		source_filepath = os.path.join(self.xcfg["source_folderpath"], f"{stock_symbol}.xlsx")
		file_exist = self.__check_file_exist(source_filepath)
		if not file_exist:
			print(f"Warning: The file {source_filepath} is not found, so skip fetching data for {stock_symbol}...")
			return
		if self.__is_excel_locked(source_filepath):
			print(f"Warning: The file {source_filepath} is locked by other process, so skip fetching data for {stock_symbol}...")
			return
		worksheet_data = self.__read_worksheet(source_filepath, expected_title_list)
		print(f"==========  {stock_symbol}  ==========")
		# import pdb; pdb.set_trace()
		if self.xcfg["target_change_pct_value"] is not None:
			_, check_change_pct_value, percentile_rank = self.analyze_change_pct_percentile(worksheet_data, expected_title_list)
			print("Change Percentage: value=%.2f%%, percentile_rank=%.2f%%" % (self.xcfg["target_change_pct_value"] * 100.0, percentile_rank))
		if self.xcfg["target_volume_value"] is not None:
			_, check_volume_value, percentile_rank = self.analyze_volume_percentile(worksheet_data, expected_title_list)
			print("Volume: value=%d, percentile_rank=%.2f%%" % (self.xcfg["target_volume_value"] / 1000, percentile_rank))
		print(f"=====================================\n")


	def show_analysis(self, is_target=False):
		if self.xcfg["stock_symbol_string"] is None:
			print("Warning: No stock to fetch...")
			return
		if is_target:
			self.__show_target_analysis()
		else:
			self.__show_analysis()


	def show_distribution_statistics(self):
		if self.xcfg["stock_symbol_string"] is None:
			print("Warning: No stock to fetch...")
			return
		expected_title_list = [self.DEFAULT_TIME_FIELD_NAME,]
		if self.xcfg["distribution_statistics"] in ["all", "volume"]:
			expected_title_list.append(self.DEFAULT_VOLUME_FIELD_NAME)
		if self.xcfg["distribution_statistics"] in ["all", "change_pct"]:
			expected_title_list.append(self.DEFAULT_EXT_CHANGE_PCT_FIELD_TITLE)
		time_index = expected_title_list.index(self.DEFAULT_TIME_FIELD_NAME)
		stock_symbol_list = self.xcfg["stock_symbol_string"].split(",")
		for stock_symbol in stock_symbol_list:
			source_filepath = os.path.join(self.xcfg["source_folderpath"], f"{stock_symbol}.xlsx")
			file_exist = self.__check_file_exist(source_filepath)
			if not file_exist:
				print(f"Warning: The file {source_filepath} is not found, so skip fetching data for {stock_symbol}...")
				continue
			if self.__is_excel_locked(source_filepath):
				print(f"Warning: The file {source_filepath} is locked by other process, so skip fetching data for {stock_symbol}...")
				continue
			# import pdb; pdb.set_trace()
			worksheet_data = self.__read_worksheet(source_filepath, expected_title_list)
			print(f"==========  {stock_symbol}  ==========")
			print("Date: %s ~ %s" % (worksheet_data["data"][0][time_index], worksheet_data["data"][-1][time_index]))
			if self.xcfg["distribution_statistics"] in ["all", "volume"]:
				volume_index = expected_title_list.index(self.DEFAULT_VOLUME_FIELD_NAME)
				volume_data_list = list(map(lambda x: x[volume_index], worksheet_data["data"]))
				volume_statistics = self.__get_statistics(volume_data_list, descending=True)
				print("Volume Statistics: %s" % volume_statistics)
			if self.xcfg["distribution_statistics"] in ["all", "change_pct"]:
				change_pct_index = expected_title_list.index(self.DEFAULT_EXT_CHANGE_PCT_FIELD_TITLE)
				change_pct_data_list = list(map(lambda x: x[change_pct_index], worksheet_data["data"]))
				change_pct_statistics = self.__get_statistics(change_pct_data_list, descending=False)
				print("Change Percentage Statistics: %s" % change_pct_statistics)
			print(f"=====================================\n")
		

if __name__ == "__main__":
# argparse 預設會把 help 文字裡的換行與多重空白「壓縮」成一行，所以你在字串裡寫的 \n 不一定會照原樣顯示。 => 建立 parser 時加上 formatter_class=argparse.RawTextHelpFormatter
	parser = argparse.ArgumentParser(description='Print help', formatter_class=argparse.RawTextHelpFormatter)
	'''
	參數基本上分兩種，一種是位置參數 (positional argument)，另一種就是選擇性參數 (optional argument)
	* example2.py
	parser.add_argument("pos1", help="positional argument 1")
	parser.add_argument("-o", "--optional-arg", help="optional argument", dest="opt", default="default")

	# python example2.py hello -o world 
	positional arg: hello
	optional arg: world
	'''
# How to add option without any argument? use action='store_true'
	'''
	'store_true' and 'store_false' - ?些是 'store_const' 分?用作存? True 和 False 值的特殊用例。
	另外，它?的默?值分?? False 和 True。例如:

	>>> parser = argparse.ArgumentParser()
	>>> parser.add_argument('--foo', action='store_true')
	>>> parser.add_argument('--bar', action='store_false')
	>>> parser.add_argument('--baz', action='store_false')
	'''
	parser.add_argument('--source_folderpath', required=False, help='Update database from the XLS files in the designated folder path. Ex: %s' % RollingStatistics.DEFAULT_DATA_FOLDERPATH)
	parser.add_argument('--print_filepath', required=False, action='store_true', help='Print the filepaths used in the process and exit.')
	parser.add_argument('--rolling_window_size', required=False, help='The number of recent trading days used for analysis. Mutually exclusive with --start_date/--end_date.')
	parser.add_argument('--start_date', required=False, help='The start date for the analysis period. Mutually exclusive with --rolling_window_size.')
	parser.add_argument('--end_date', required=False, help='The end date for the analysis period. Mutually exclusive with --rolling_window_size.')
	parser.add_argument('-s', '--show_analysis', required=False, action='store_true', help='Show the analysis results and exit.')
	parser.add_argument('--show_distribution_statistics', required=False, action='store_true', help='Show the distribution statistics of the data and exit.')
	parser.add_argument(
		'distribution_statistics', 
		type=parse_distribution_statistics, 
		nargs='?',  #　後面的值可以有，也可以沒有。
		const='all', #　如果後面沒有值，就使用 const 的值。
		default='all',
		help='''The statistics to show. 
			'0/all: all
			 1/change_pct: change percentage 
			 2/volume: volume.''')
	parser.add_argument('--stock_symbol_list', required=False, help='The stock symbol list. Multiple stock symbols are seperated by comma. Ex: 00850.TW,00881.TW,00692.TW,MSFT,GOOG')
	parser.add_argument('--check_date', required=False, help='The date to check. Ex: 2023-01-01. If not specified, the latest date in the data will be used.')
	parser.add_argument('--target_stock_symbol', required=False, help='The target stock symbol. Mutually exclusive with --stock_symbol_list.')
	parser.add_argument('--target_change_pct_value', required=False, help='The target change percentage to check. Ex: 3.11. Mutually exclusive with --stock_symbol_list.')
	parser.add_argument('--target_volume_value', required=False, help='The target volume value to check. Ex: 14200. Mutually exclusive with --stock_symbol_list.')

	args = parser.parse_args()
	# import pdb; pdb.set_trace()
	cfg = {}
	is_target = False
	if args.source_folderpath is not None: cfg['source_folderpath'] = args.source_folderpath
	if args.rolling_window_size is not None: cfg['rolling_window_size'] = int(args.rolling_window_size)
	if args.rolling_window_size is None:
		if args.start_date is not None: cfg['start_date'] = args.start_date
		if args.end_date is not None: cfg['end_date'] = args.end_date
	if args.distribution_statistics is not None: cfg['distribution_statistics'] = args.distribution_statistics
	if args.stock_symbol_list is not None: cfg['stock_symbol_string'] = args.stock_symbol_list
	if args.check_date is not None: cfg['check_date'] = args.check_date
	if args.stock_symbol_list is None:
		is_target = True
		if args.target_stock_symbol is not None: cfg['stock_symbol_string'] = args.target_stock_symbol
		if args.target_change_pct_value is not None: cfg['target_change_pct_value'] = float(args.target_change_pct_value) / 100.0
		if args.target_volume_value is not None: cfg['target_volume_value'] = int(args.target_volume_value) * 1000
	# import pdb; pdb.set_trace()
	with RollingStatistics(cfg) as obj:
		if args.print_filepath:
			obj.print_filepath()
			sys.exit(0)
		if args.show_analysis:
			obj.show_analysis(is_target)
			sys.exit(0)
		if args.show_distribution_statistics:
			obj.show_distribution_statistics()
			sys.exit(0)